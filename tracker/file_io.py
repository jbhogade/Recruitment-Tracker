# ==========================================================================
# File : tracker/file_io.py
# What : Reads and writes the three file types used by Import, Export and
#        Template:   CSV (.csv),  Excel (.xlsx),  Excel 97-2003 (.xls)
#        This file knows NOTHING about candidates; it only handles tables.
# ==========================================================================

#-------------- Start Code for Imports and file type list -----------------
import csv
import io
from datetime import date, datetime

# key used in the web address (?format=...) -> details of that file type
FORMATS = {
    "csv":  {"label": "CSV (.csv)",                     "ext": "csv",  "type": "text/csv; charset=utf-8"},
    "xlsx": {"label": "Excel (.xlsx)",                  "ext": "xlsx", "type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    "xls":  {"label": "Excel 97-2003 (.xls)",           "ext": "xls",  "type": "application/vnd.ms-excel"},
}

SHEET_NAME = "Candidates"
DATE_TEXT = "%d-%b-%Y"            # how dates appear in CSV files, e.g. 25-Oct-2026
DATE_EXCEL = "DD-MMM-YYYY"        # how dates look in Excel cells


class FileReadError(Exception):
    """Raised with a plain-language message when a file cannot be read."""
#-------------- End of the code ------------------------------


#-------------- Start Code for Protecting text from Excel formulas -----------------
def safe_text(value):
    """Text starting with = + - @ could run as a formula when opened in Excel.
    A leading apostrophe makes Excel treat it as plain text. Import removes it again."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value
#-------------- End of the code ------------------------------


#-------------- Start Code for Reading a file (CSV / XLS / XLSX) -----------------
def detect_type(content, filename):
    """Decide the real file type from the first bytes (not only from the name)."""
    name = filename.lower()
    if content[:2] == b"PK":
        return "xlsx"
    if content[:4] == b"\xd0\xcf\x11\xe0":
        return "xls"
    if name.endswith(".csv"):
        return "csv"
    if name.endswith((".xlsx", ".xls", ".xlsm")):
        raise FileReadError("This file looks damaged or is not a real Excel file. Open it in Excel and save it again as .xlsx or .csv.")
    raise FileReadError("Unsupported file type. Please upload a .csv, .xlsx or .xls file.")


def read_table(content, filename):
    """Return (headings, rows).
    headings : list of heading text from the first non-empty row.
    rows     : list of (row number in the file, list of cell values), empty rows skipped.
    """
    kind = detect_type(content, filename)
    try:
        if kind == "xlsx":
            table = _read_xlsx(content)
        elif kind == "xls":
            table = _read_xls(content)
        else:
            table = _read_csv(content)
    except FileReadError:
        raise
    except Exception as error:     # damaged file, password protected, etc.
        raise FileReadError("The file could not be read (it may be damaged or password protected). "
                            f"Technical detail: {type(error).__name__}.") from error

    numbered = [(number, list(values)) for number, values in table
                if any(str(v).strip() for v in values if v is not None)]
    if not numbered:
        raise FileReadError("The file is empty.")
    headings = [str(v).strip() if v is not None else "" for v in numbered[0][1]]
    return headings, numbered[1:]


def _read_csv(content):
    """CSV: tries UTF-8 (with or without BOM), then Windows-1252 (older Excel files)."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise FileReadError("The CSV file text encoding is not supported. Save it as 'CSV UTF-8' from Excel.")
    first_line = text.splitlines()[0] if text.strip() else ""
    delimiter = max(",;\t", key=first_line.count) if first_line else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [(number, row) for number, row in enumerate(reader, start=1)]


def _read_xlsx(content):
    """Excel .xlsx: uses the sheet named 'Candidates', otherwise the first sheet."""
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = book[SHEET_NAME] if SHEET_NAME in book.sheetnames else book.worksheets[0]
    rows = [(number, row) for number, row in enumerate(sheet.iter_rows(values_only=True), start=1)]
    book.close()
    return rows


def _read_xls(content):
    """Excel 97-2003 .xls: uses the sheet named 'Candidates', otherwise the first sheet."""
    import xlrd
    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_name(SHEET_NAME) if SHEET_NAME in book.sheet_names() else book.sheet_by_index(0)
    rows = []
    for index in range(sheet.nrows):
        values = []
        for cell in sheet.row(index):
            if cell.ctype == xlrd.XL_CELL_DATE:
                values.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                values.append(None)
            elif cell.ctype == xlrd.XL_CELL_ERROR:
                values.append(None)
            else:
                values.append(cell.value)
        rows.append((index + 1, values))
    return rows
#-------------- End of the code ------------------------------


#-------------- Start Code for Writing a file (CSV / XLS / XLSX) -----------------
def write_table(kind, headings, rows, date_columns=(), text_columns=(), number_columns=(),
                notes=None, dropdowns=None, sheet_name=SHEET_NAME):
    """Create the file and return its bytes.

    headings        : list of column headings
    rows            : list of lists (values may be text, numbers, date or None)
    date_columns    : column positions (0 = first) that hold dates
    text_columns    : column positions kept as text (e.g. phone numbers)
    number_columns  : column positions with decimal numbers
    notes           : optional list of rows for an extra "Instructions" sheet (Excel files only)
    dropdowns       : {column position: [allowed values]} drop-down lists (.xlsx only)
    """
    if kind == "csv":
        return _write_csv(headings, rows)
    if kind == "xlsx":
        return _write_xlsx(headings, rows, date_columns, text_columns, number_columns, notes, dropdowns, sheet_name)
    if kind == "xls":
        return _write_xls(headings, rows, date_columns, text_columns, number_columns, notes, sheet_name)
    raise ValueError("Unknown file type")


def _write_csv(headings, rows):
    """CSV with a BOM so Excel shows every character correctly."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headings)
    for row in rows:
        writer.writerow([
            value.strftime(DATE_TEXT) if isinstance(value, (date, datetime)) else
            ("" if value is None else safe_text(value))
            for value in row
        ])
    return buffer.getvalue().encode("utf-8-sig")


def _write_xlsx(headings, rows, date_columns, text_columns, number_columns, notes, dropdowns, sheet_name):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.append(headings)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5D8A")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append([safe_text(v) for v in row])

    last_row = max(len(rows) + 1, 501) if dropdowns else len(rows) + 1
    for position in date_columns:
        for r in range(2, last_row + 1):
            sheet.cell(row=r, column=position + 1).number_format = DATE_EXCEL
    for position in text_columns:
        for r in range(2, last_row + 1):
            sheet.cell(row=r, column=position + 1).number_format = "@"
    for position in number_columns:
        for r in range(2, last_row + 1):
            sheet.cell(row=r, column=position + 1).number_format = "0.0#"
    for position, options in (dropdowns or {}).items():
        letter = get_column_letter(position + 1)
        rule = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        sheet.add_data_validation(rule)
        rule.add(f"{letter}2:{letter}{last_row}")

    for position, heading in enumerate(headings, start=1):
        longest = max([len(str(heading))] + [len(str(r[position - 1])) for r in rows[:200] if position - 1 < len(r)])
        sheet.column_dimensions[get_column_letter(position)].width = min(max(longest + 3, 12), 45)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headings))}{max(len(rows) + 1, 2)}"

    if notes:
        info = book.create_sheet("Instructions")
        for row in notes:
            info.append(list(row))
        for cell in info[1]:
            cell.font = Font(bold=True)
        for index, width in enumerate((26, 14, 70, 28), start=1):
            info.column_dimensions[get_column_letter(index)].width = width
        for row in info.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _write_xls(headings, rows, date_columns, text_columns, number_columns, notes, sheet_name):
    import xlwt

    book = xlwt.Workbook(encoding="utf-8")
    sheet = book.add_sheet(sheet_name)
    head_style = xlwt.easyxf("font: bold on, colour white; pattern: pattern solid, fore_colour dark_blue;")
    date_style = xlwt.easyxf(num_format_str=DATE_EXCEL)
    text_style = xlwt.easyxf(num_format_str="@")
    number_style = xlwt.easyxf(num_format_str="0.0#")

    for column, heading in enumerate(headings):
        sheet.write(0, column, heading, head_style)
        longest = max([len(str(heading))] + [len(str(r[column])) for r in rows[:200] if column < len(r)])
        sheet.col(column).width = 256 * min(max(longest + 3, 12), 45)
    for r, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            if value is None or value == "":
                continue
            if column in date_columns and isinstance(value, (date, datetime)):
                sheet.write(r, column, value, date_style)
            elif column in number_columns and isinstance(value, (int, float)):
                sheet.write(r, column, value, number_style)
            elif column in text_columns:
                sheet.write(r, column, str(safe_text(value)), text_style)
            else:
                sheet.write(r, column, safe_text(value))
    sheet.set_panes_frozen(True)
    sheet.set_horz_split_pos(1)

    if notes:
        info = book.add_sheet("Instructions")
        bold = xlwt.easyxf("font: bold on; align: wrap on, vert top;")
        wrap = xlwt.easyxf("align: wrap on, vert top;")
        for r, row in enumerate(notes):
            for column, value in enumerate(row):
                info.write(r, column, value, bold if r == 0 else wrap)
        for index, width in enumerate((26, 14, 70, 28)):
            info.col(index).width = 256 * width

    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
#-------------- End of the code ------------------------------
